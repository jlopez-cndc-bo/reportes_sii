"""Generación de archivos Excel (.xlsx) a partir del resultado de un reporte.

Reproduce el estilo del export de la app original: título combinado, encabezado
resaltado, paneles congelados y formatos numéricos/fecha por columna.
"""
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="1E4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri")
_TITLE_FONT = Font(bold=True, size=13, name="Arial Narrow")


def _num_format(fmt: str | None) -> str | None:
    if fmt == "datetime":
        return "dd/mm/yyyy hh:mm"
    if fmt and fmt.startswith("0."):
        return "#,##0." + "0" * (len(fmt.split(".")[1]))
    return None


def _llenar_hoja(ws, result: dict[str, Any], titulo: str) -> None:
    cols = [c for c in result["columns"] if not c.get("hidden")]
    ncols = len(cols)
    # Fila 1: título combinado.
    ws.cell(row=1, column=1, value=titulo).font = _TITLE_FONT
    if ncols > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)

    aviso = result.get("aviso")
    header_row = 3 if aviso else 2
    if aviso:
        ws.cell(row=2, column=1, value=f"Nota: {aviso}").font = Font(italic=True, color="B45309")
        if ncols > 1:
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)

    # Encabezado.
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=header_row, column=j, value=c["label"])
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # Datos.
    fmt_by_col = [_num_format(c.get("format")) for c in cols]
    r = header_row + 1
    for row in result["rows"]:
        for j, c in enumerate(cols, start=1):
            v = row.get(c["key"])
            if c.get("format") == "datetime" and isinstance(v, str):
                try:
                    v = datetime.fromisoformat(v)
                except ValueError:
                    pass
            cell = ws.cell(row=r, column=j, value=v)
            nf = fmt_by_col[j - 1]
            if nf and v is not None:
                cell.number_format = nf
        r += 1

    # Anchos aproximados.
    for j, c in enumerate(cols, start=1):
        width = max(10, min(40, len(str(c["label"])) + 2))
        ws.column_dimensions[get_column_letter(j)].width = width

    # Congelar encabezado + primera columna.
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)


def build_workbook(result: dict[str, Any], titulo: str) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = (result.get("table") or "Reporte")[:31]
    _llenar_hoja(ws, result, titulo)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_workbook_multi(hojas: list[tuple[str, dict[str, Any], str]]) -> BytesIO:
    """Un workbook con una hoja por elemento de `hojas`: (nombre, result, titulo)."""
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, result, titulo in hojas:
        ws = wb.create_sheet(nombre[:31])
        _llenar_hoja(ws, result, titulo)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
