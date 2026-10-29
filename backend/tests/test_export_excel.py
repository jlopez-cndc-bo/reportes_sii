"""Tests de generación de workbooks (sin BD)."""
from openpyxl import load_workbook

from app.export_excel import build_workbook, build_workbook_multi

_RESULT = {
    "table": "DEMO",
    "columns": [{"key": "a", "label": "A"}, {"key": "b", "label": "B", "format": "0.0"}],
    "rows": [{"a": "x", "b": 1.5}, {"a": "y", "b": 2.25}],
}


def test_build_workbook_hoja_unica():
    buf = build_workbook(_RESULT, "Título")
    wb = load_workbook(buf)
    assert wb.sheetnames == ["DEMO"]
    ws = wb["DEMO"]
    assert ws.cell(row=1, column=1).value == "Título"
    assert ws.cell(row=2, column=1).value == "A"
    assert ws.cell(row=3, column=1).value == "x"


def test_build_workbook_multi_una_hoja_por_direccion():
    buf = build_workbook_multi([
        ("Inyección", _RESULT, "Título Iny"),
        ("Retiro", _RESULT, "Título Ret"),
    ])
    wb = load_workbook(buf)
    assert wb.sheetnames == ["Inyección", "Retiro"]
    assert wb["Inyección"].cell(row=1, column=1).value == "Título Iny"
    assert wb["Retiro"].cell(row=1, column=1).value == "Título Ret"
    # Cada hoja arranca de cero (no hereda la hoja 'Sheet' por defecto).
    assert wb["Inyección"].cell(row=3, column=1).value == "x"
